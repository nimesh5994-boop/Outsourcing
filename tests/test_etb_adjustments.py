"""End-to-end verification of the Journals sheet (see excel_builder.
build_journals_sheet_formulas and data_sheets.JournalsRef) - the single
place a preparer posts a narrated adjusting journal, replacing what used
to be a flat, unnamed Adjustment cell typed directly onto the TB Tie-Out
sheet.

Builds a real workbook through the actual pipeline (recon.run_all_recons +
tb_tieout.build_tieout, same as main.py's generation step), types a
journal row into the Journals sheet, and uses the `formulas` library to
*evaluate* the result - not just eyeball the formula strings - confirming
the journal reaches DATA_TB_Current's balance and, through it, the P&L
statement, with no changes needed in any of those sheets' own formulas.
"""
from pathlib import Path

import openpyxl
import pytest

pytest.importorskip("formulas")
import formulas  # noqa: E402

from app import control_accounts, corporation_tax, fixed_assets, nominal_matrix, recon, tb_tieout  # noqa: E402
from app.excel_builder import build_workbook  # noqa: E402


def _evaluate(path) -> dict:
    xl = formulas.ExcelModel().loads(str(path)).finish()
    return xl.calculate()


def _cell(sol: dict, workbook_filename: str, sheet: str, coord: str):
    key = f"'[{workbook_filename}]{sheet.upper()}'!{coord}"
    return sol[key].value[0][0]


def _build_and_save(canonical_data, out_path):
    results = recon.run_all_recons(canonical_data)
    results = results + [tb_tieout.build_tieout(
        canonical_data.get("tb_current"), canonical_data.get("tb_comparative"), canonical_data.get("nominal_current"),
    )]
    ca_results = control_accounts.build_all_rollforwards(
        canonical_data["tb_current"], canonical_data["tb_comparative"], canonical_data["nominal_current"],
        canonical_data["aged_debtors"], canonical_data["aged_creditors"],
    )
    mx_results = nominal_matrix.build_all_matrices(canonical_data["tb_current"], canonical_data["nominal_current"])
    ct_computation = corporation_tax.compute(accounting_profit=float(canonical_data["pl_current"]["amount"].sum()))
    fixed_asset_result = fixed_assets.category_level_rollforward(
        canonical_data["tb_current"], canonical_data["tb_comparative"], canonical_data["nominal_current"]
    )
    wb = build_workbook(
        "Brightwell Landscaping Supplies Limited",
        "Year ended 31 December 2025", "Year ended 31 December 2024",
        canonical_data, results, control_account_results=ca_results, matrix_results=mx_results,
        ct_computation=ct_computation, fixed_asset_result=fixed_asset_result,
    )
    wb.save(out_path)
    return wb


def _journals_sheet(wb):
    name = next(s for s in wb.sheetnames if s.endswith("Journals"))
    ws = wb[name]
    header_row = next(r for r in range(1, 10) if ws.cell(row=r, column=1).value == "No.")
    return name, ws, header_row + 1  # sheet name, worksheet, first typeable row


def test_journals_sheet_exists_and_is_genuinely_blank(tmp_path, canonical_data):
    out = tmp_path / "wp.xlsx"
    wb = _build_and_save(canonical_data, out)
    name, ws, first_row = _journals_sheet(wb)
    assert [ws.cell(row=first_row - 1, column=c).value for c in range(1, 6)] == ["No.", "Narration", "Account Code", "Debit", "Credit"]
    assert ws.cell(row=first_row, column=1).value is None  # genuinely blank, not a formula
    assert ws.cell(row=first_row, column=4).value is None


def test_posting_a_journal_flows_through_to_data_tb_current_and_the_pl_statement(tmp_path, canonical_data):
    out = tmp_path / "wp.xlsx"
    wb = _build_and_save(canonical_data, out)
    _, ws, first_row = _journals_sheet(wb)
    pl_sheet_name = next(s for s in wb.sheetnames if s.endswith("Profit and Loss"))

    pl_code = str(canonical_data["pl_current"].iloc[0]["account_code"])
    dtc_row = _dtc_row_for(wb, pl_code)

    sol_before = _evaluate(out)
    before_dtc = _cell(sol_before, "wp.xlsx", "DATA_TB_Current", f"F{dtc_row}")
    before_net_profit = _cell(sol_before, "wp.xlsx", pl_sheet_name, "B12")

    JOURNAL_AMOUNT = 250.0
    ws.cell(row=first_row, column=1, value="JNL 1")
    ws.cell(row=first_row, column=2, value="Test correcting entry")
    ws.cell(row=first_row, column=3, value=pl_code)
    ws.cell(row=first_row, column=4, value=JOURNAL_AMOUNT)  # Debit
    wb.save(out)

    sol_after = _evaluate(out)
    after_dtc = _cell(sol_after, "wp.xlsx", "DATA_TB_Current", f"F{dtc_row}")
    after_net_profit = _cell(sol_after, "wp.xlsx", pl_sheet_name, "B12")

    assert after_dtc == pytest.approx(before_dtc + JOURNAL_AMOUNT)
    # DATA_PL negates DATA_TB_Current's balance (see data_sheets._write_derived_amount_sheet),
    # so a +250 debit journal should move Net Profit by -250.
    assert after_net_profit == pytest.approx(before_net_profit - JOURNAL_AMOUNT)


def test_journals_balance_check_flags_an_unbalanced_register(tmp_path, canonical_data):
    out = tmp_path / "wp.xlsx"
    wb = _build_and_save(canonical_data, out)
    name, ws, first_row = _journals_sheet(wb)

    pl_code = str(canonical_data["pl_current"].iloc[0]["account_code"])
    ws.cell(row=first_row, column=1, value="JNL 1")
    ws.cell(row=first_row, column=3, value=pl_code)
    ws.cell(row=first_row, column=4, value=100.0)  # a lone debit, no matching credit anywhere
    wb.save(out)

    sol = _evaluate(out)
    # the balance-check cell sits two rows below the last typeable row
    check_row = first_row + 200 + 1  # JOURNALS_DATA_ROWS = 200
    flag = _cell(sol, "wp.xlsx", name, f"D{check_row}")
    assert flag == "OUT OF BALANCE - CHECK ENTRIES"


def _dtc_row_for(wb, account_code: str) -> int:
    ws = wb["DATA_TB_Current"]
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == account_code:
            return r
    raise AssertionError(f"account code {account_code} not found on DATA_TB_Current")
