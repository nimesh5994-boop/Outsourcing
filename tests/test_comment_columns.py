"""Every place a preparer explains an exception in this workbook (TB
Tie-Out's opening balance disagreements, Loan/Stock reviews, P&L Variance
Review, TB Lead Schedule) uses the same "Comment" convention: a genuinely
blank cell filled with the pale-yellow "type here" convention
(excel_builder.INPUT_FILL). Covers the two mechanisms that produce it:
excel_builder._write_dataframe (any column literally named "Comment")
and TB Lead Schedule's own hand-written column.
"""
import pandas as pd
from openpyxl import Workbook

from app.excel_builder import INPUT_FILL, _write_dataframe, build_tb_lead_schedule_formulas
from app.data_sheets import write_data_sheets


def test_write_dataframe_fills_a_comment_column_but_not_others():
    wb = Workbook()
    ws = wb.active
    df = pd.DataFrame([{"Account": "1000", "Amount": 500.0, "Comment": ""}])
    _write_dataframe(ws, df, start_row=1)

    assert ws.cell(row=2, column=1).fill.fgColor.rgb != INPUT_FILL.fgColor.rgb  # Account
    assert ws.cell(row=2, column=2).fill.fgColor.rgb != INPUT_FILL.fgColor.rgb  # Amount
    assert ws.cell(row=2, column=3).fill.fgColor.rgb == INPUT_FILL.fgColor.rgb  # Comment


def test_tb_lead_schedule_review_column_is_input_styled():
    tb_current = pd.DataFrame([{"account_code": "200", "account_name": "Sales", "account_type": "Revenue", "debit": 0.0, "credit": 1000.0, "balance": -1000.0}])
    variance_detail = pd.DataFrame([{"account_code": "200", "account_name": "Sales", "current_year": -1000.0, "comparative_year": 0.0, "variance_amount": -1000.0, "variance_pct": 1.0, "flag": True}])

    wb = Workbook()
    refs = write_data_sheets(wb, {"tb_current": tb_current})
    build_tb_lead_schedule_formulas(wb, "Test Client", "Year ended 31 December 2025", "1", variance_detail, refs)
    ws = wb["1 TB Lead Schedule"]

    header_row = next(r for r in range(1, 10) if ws.cell(row=r, column=8).value == "Reviewed / reallocation required?")
    data_row = header_row + 1
    assert ws.cell(row=data_row, column=8).value == ""
    assert ws.cell(row=data_row, column=8).fill.fgColor.rgb == INPUT_FILL.fgColor.rgb
