"""Xero connection + pull, end to end through the real HTTP routes and
Postgres, against a fake Xero server (no network, no credentials): OAuth
connect, pulling the TB/ledger into a job, that pulled data flowing through
the full generate pipeline into the workbook, and that an uploaded export
still overrides a pulled report. The API JSON the fake serves is Xero's
documented shape as understood when written - see test_xero_api.py's note
on checking the first live pull against the raw report stored with it."""
import io
import json
import urllib.parse
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from app import parsers, xero_api, xero_reports
from tests.test_storage_and_routes import _signup, http_client, pytestmark  # noqa: F401  (fixture + DB skip marker)
from tests.test_xero_api import MS_2025_06_15, _journal, _line, _tb_report

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_data"
TOKEN_SCOPES = "offline_access accounting.settings accounting.reports.trialbalance.read"


def _sample_tb():
    cur, comp = xero_reports.parse_trial_balance(parsers.FileDataSource(SAMPLE_DIR / "trial_balance_current_xero.xlsx"))
    return cur, comp


def _report_and_accounts(df: pd.DataFrame, id_prefix: str):
    accounts = {f"{id_prefix}{i}": {"code": r.account_code, "name": r.account_name, "type": r.account_type} for i, r in enumerate(df.itertuples())}
    rows = [(f"{r.account_name} ({r.account_code})", f"{id_prefix}{i}", r.debit, r.credit) for i, r in enumerate(df.itertuples())]
    return _tb_report({"All": rows}), accounts


_TYPE_TOKEN = {v: k for k, v in xero_api.ACCOUNT_TYPE_LABELS.items() if k in ("BANK", "CURRENT", "CURRLIAB", "DIRECTCOSTS", "EQUITY", "FIXED", "LIABILITY", "OVERHEADS", "SALES", "REVENUE")}


class FakeXero:
    def __init__(self, scopes=TOKEN_SCOPES, tenants=None, journals=None, forbid_journals=False, tb_comparative_shift=0.0):
        self.scopes = scopes
        self.tenants = tenants or [{"id": "conn-1", "tenantId": "tenant-1", "tenantName": "Brightwell Landscaping (Xero)", "tenantType": "ORGANISATION"}]
        self.journals, self.forbid_journals = journals or [], forbid_journals
        self.calls = []
        cur, comp = _sample_tb()
        self.cur_report, self.accounts = _report_and_accounts(cur, "a")
        self.comp_report, _ = _report_and_accounts(comp, "a")
        self.tb_current = cur

    def __call__(self, method, url, headers, body=None):
        self.calls.append((method, url))
        parsed = urllib.parse.urlparse(url)
        if parsed.netloc == "identity.xero.com":
            return 200, {}, json.dumps({"access_token": "ACCESS-1", "refresh_token": "REFRESH-1", "expires_in": 1800, "scope": self.scopes}).encode()
        if parsed.netloc == "api.xero.com" and parsed.path.startswith("/connections"):
            return (204, {}, b"") if method == "DELETE" else (200, {}, json.dumps(self.tenants).encode())
        path = parsed.path.split("/api.xro/2.0/", 1)[1]
        q = urllib.parse.parse_qs(parsed.query)
        assert headers.get("xero-tenant-id") == "tenant-1"
        if path == "Accounts":
            return 200, {}, json.dumps({"Accounts": [
                {"AccountID": k, "Code": v["code"], "Name": v["name"], "Type": _TYPE_TOKEN.get(v["type"], "CURRENT")} for k, v in self.accounts.items()
            ]}).encode()
        if path == "Organisation":
            return 200, {}, json.dumps({"Organisations": [{"Name": "Brightwell", "FinancialYearEndDay": 31, "FinancialYearEndMonth": 12}]}).encode()
        if path == "Reports/TrialBalance":
            report = self.cur_report if q["date"][0] == "2025-12-31" else self.comp_report
            return 200, {}, json.dumps(report).encode()
        if path == "Journals":
            if self.forbid_journals:
                return 403, {}, b'{"Detail":"AuthorizationUnsuccessful"}'
            offset = int(q.get("offset", ["0"])[0])
            page = [j for j in self.journals if j["JournalNumber"] > offset][:100]
            return 200, {}, json.dumps({"Journals": page}).encode()
        raise AssertionError(f"unexpected Xero call {method} {url}")


@pytest.fixture
def xero_env(monkeypatch):
    monkeypatch.setenv("XERO_CLIENT_ID", "cid")
    monkeypatch.setenv("XERO_CLIENT_SECRET", "csecret")
    monkeypatch.setenv("XERO_REDIRECT_URI", "https://wp.test/xero/callback")
    monkeypatch.delenv("XERO_SCOPES", raising=False)


def _new_client_and_job(c, email):
    practice_id = _signup(c, admin_email=email)
    resp = c.post(f"/practices/{practice_id}/clients", data={"name": "Brightwell"}, follow_redirects=False)
    client_id = resp.headers["location"].rsplit("/", 1)[-1]
    resp = c.post(f"/clients/{client_id}/jobs", data={
        "current_period_start": "2025-01-01", "current_period_end": "2025-12-31",
        "comparative_period_start": "2024-01-01", "comparative_period_end": "2024-12-31",
    }, follow_redirects=False)
    return client_id, resp.headers["location"].rsplit("/", 1)[-1]


def _connect(c, client_id):
    resp = c.get(f"/clients/{client_id}/xero/connect", follow_redirects=False)
    assert resp.status_code == 303
    location = resp.headers["location"]
    assert location.startswith("https://login.xero.com/identity/connect/authorize?")
    q = urllib.parse.parse_qs(urllib.parse.urlparse(location).query)
    assert q["redirect_uri"] == ["https://wp.test/xero/callback"]
    return c.get(f"/xero/callback?code=abc&state={q['state'][0]}", follow_redirects=False)


def test_connect_pull_generate_puts_the_pulled_trial_balance_in_the_workbook(http_client, xero_env, monkeypatch):
    c = http_client
    fake = FakeXero()
    monkeypatch.setattr(xero_api, "default_http", fake)
    client_id, job_id = _new_client_and_job(c, "xero-e2e@acme.test")

    resp = _connect(c, client_id)
    assert resp.headers["location"] == f"/clients/{client_id}?xero=connected"

    from app import storage
    conn = storage.get_xero_connection(client_id)
    assert conn["tenant_id"] == "tenant-1" and conn["tenant_name"].startswith("Brightwell")
    assert "ACCESS-1" not in json.dumps(conn) and "REFRESH-1" not in json.dumps(conn)  # credentials are encrypted at rest

    resp = c.post(f"/jobs/{job_id}/xero/pull", follow_redirects=False)
    assert resp.status_code == 303
    job = storage.get_job(job_id)
    assert job["xero_pull"]["summary"]["tb_current_rows"] == len(fake.tb_current)
    assert job["xero_pull"]["journals"]["status"] == "unavailable"  # token has no journals scope
    assert not job["xero_pull"]["warnings"]

    assert c.post(f"/jobs/{job_id}/generate", follow_redirects=False).status_code == 303
    assert storage.get_job(job_id)["status"] == "generated"
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/jobs/{job_id}/download").content))
    tb_sheet = wb["DATA_TB_Current"]
    codes = {str(r[0]) for r in tb_sheet.iter_rows(min_row=2, values_only=True) if r[0] is not None}
    assert codes == set(fake.tb_current["account_code"].astype(str))
    assert any("ETB" in name for name in wb.sheetnames)
    summary = {s["name"]: s for s in storage.get_job(job_id)["summary"]}
    assert summary["TB self-balance check"]["status"] == "ok"


def test_job_and_client_pages_render_with_xero_states(http_client, xero_env, monkeypatch):
    c = http_client
    monkeypatch.setattr(xero_api, "default_http", FakeXero())
    client_id, job_id = _new_client_and_job(c, "xero-pages@acme.test")
    assert "Connect to Xero" in c.get(f"/clients/{client_id}").text
    assert "Connect it on the client page" in c.get(f"/jobs/{job_id}").text
    _connect(c, client_id)
    assert "Connected to" in c.get(f"/clients/{client_id}").text
    assert "Pull Trial Balance" in c.get(f"/jobs/{job_id}").text
    c.post(f"/jobs/{job_id}/xero/pull")
    page = c.get(f"/jobs/{job_id}").text
    assert "Re-pull Trial Balance" in page and "accounts" in page


def test_unconfigured_server_redirects_with_a_clear_message(http_client, monkeypatch):
    monkeypatch.delenv("XERO_CLIENT_ID", raising=False)
    monkeypatch.delenv("XERO_CLIENT_SECRET", raising=False)
    client_id, _job = _new_client_and_job(http_client, "xero-unconf@acme.test")
    resp = http_client.get(f"/clients/{client_id}/xero/connect", follow_redirects=False)
    assert resp.headers["location"] == f"/clients/{client_id}?xero=not_configured"
    assert "isn't set up on this server" in http_client.get(f"/clients/{client_id}").text


def test_callback_rejects_a_tampered_state(http_client, xero_env, monkeypatch):
    monkeypatch.setattr(xero_api, "default_http", FakeXero())
    _new_client_and_job(http_client, "xero-state@acme.test")
    assert http_client.get("/xero/callback?code=abc&state=forged", follow_redirects=False).status_code == 400


def test_denied_consent_connects_nothing(http_client, xero_env, monkeypatch):
    monkeypatch.setattr(xero_api, "default_http", FakeXero())
    client_id, _job = _new_client_and_job(http_client, "xero-denied@acme.test")
    state = urllib.parse.parse_qs(urllib.parse.urlparse(http_client.get(f"/clients/{client_id}/xero/connect", follow_redirects=False).headers["location"]).query)["state"][0]
    resp = http_client.get(f"/xero/callback?error=access_denied&state={state}", follow_redirects=False)
    assert resp.headers["location"] == f"/clients/{client_id}?xero=denied"
    from app import storage
    assert storage.get_xero_connection(client_id) is None


def test_multiple_organisations_need_a_choice_before_a_pull(http_client, xero_env, monkeypatch):
    tenants = [{"id": "c1", "tenantId": "tenant-1", "tenantName": "Org One", "tenantType": "ORGANISATION"},
               {"id": "c2", "tenantId": "tenant-2", "tenantName": "Org Two", "tenantType": "ORGANISATION"}]
    monkeypatch.setattr(xero_api, "default_http", FakeXero(tenants=tenants))
    client_id, job_id = _new_client_and_job(http_client, "xero-multi@acme.test")
    assert _connect(http_client, client_id).headers["location"].endswith("?xero=choose")
    from app import storage
    assert not storage.get_xero_connection(client_id).get("tenant_id")
    http_client.post(f"/jobs/{job_id}/xero/pull")
    assert "isn't connected to a Xero organisation" in storage.get_job(job_id)["xero_message"]
    http_client.post(f"/clients/{client_id}/xero/tenant", data={"tenant_id": "tenant-1"})
    assert storage.get_xero_connection(client_id)["tenant_name"] == "Org One"
    assert http_client.post(f"/clients/{client_id}/xero/tenant", data={"tenant_id": "nope"}).status_code == 400


def test_disconnect_removes_stored_credentials(http_client, xero_env, monkeypatch):
    fake = FakeXero()
    monkeypatch.setattr(xero_api, "default_http", fake)
    client_id, _job = _new_client_and_job(http_client, "xero-disc@acme.test")
    _connect(http_client, client_id)
    http_client.post(f"/clients/{client_id}/xero/disconnect")
    from app import storage
    assert storage.get_xero_connection(client_id) is None
    assert any(m == "DELETE" for m, _ in fake.calls)  # revoked on Xero's side too


def test_part_year_job_surfaces_a_financial_year_warning(http_client, xero_env, monkeypatch):
    monkeypatch.setattr(xero_api, "default_http", FakeXero())
    practice_id = _signup(http_client, admin_email="xero-fy@acme.test")
    client_id = http_client.post(f"/practices/{practice_id}/clients", data={"name": "Part Year"}, follow_redirects=False).headers["location"].rsplit("/", 1)[-1]
    job_id = http_client.post(f"/clients/{client_id}/jobs", data={
        "current_period_start": "2025-07-01", "current_period_end": "2025-12-31",
        "comparative_period_start": "", "comparative_period_end": "",
    }, follow_redirects=False).headers["location"].rsplit("/", 1)[-1]
    _connect(http_client, client_id)
    http_client.post(f"/jobs/{job_id}/xero/pull")
    from app import storage
    warnings = storage.get_job(job_id)["xero_pull"]["warnings"]
    assert warnings and "financial year ends 31 December" in warnings[0]


def test_an_uploaded_trial_balance_overrides_the_pulled_one(http_client, xero_env, monkeypatch):
    c = http_client
    fake = FakeXero()
    # make the pulled TB visibly different from the uploaded sample on one account
    first = fake.cur_report["Reports"][0]["Rows"][1]["Rows"][0]["Cells"]
    first[3]["Value"], first[4]["Value"] = "777.00", "0.00"
    monkeypatch.setattr(xero_api, "default_http", fake)
    client_id, job_id = _new_client_and_job(c, "xero-override@acme.test")
    _connect(c, client_id)
    c.post(f"/jobs/{job_id}/xero/pull")

    from app import main, storage
    pulled = main._load_canonical_data(storage.get_job(job_id))["tb_current"]
    assert pulled.iloc[0]["debit"] == 777.0

    resp = c.post(f"/jobs/{job_id}/uploads", data={"section_report_type": "trial_balance"},
                  files={"files": ("tb.xlsx", (SAMPLE_DIR / "trial_balance_current_xero.xlsx").read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                  follow_redirects=False)
    assert resp.status_code == 303
    job = storage.get_job(job_id)
    for upload in job["uploads"].values():
        upload["confirmed"] = True
        upload["period"] = "current"
    storage.save_job(job)
    uploaded, _ = _sample_tb()
    assert main._load_canonical_data(storage.get_job(job_id))["tb_current"].iloc[0]["debit"] == uploaded.iloc[0]["debit"]


# ---------- ledger pull (Journals), resumable ----------

def _many_journals():
    out = []
    for n in range(1, 251):
        ms = 1700000000000 if n <= 120 else MS_2025_06_15  # first 120 are before the 2025 period
        out.append({**_journal(ms, "MANJOURNAL", [_line("200", "Sales", -float(n)), _line("610", "Debtors", float(n))]), "JournalNumber": n})
    return out


def test_ledger_pull_resumes_across_batches_and_keeps_only_period_journals(http_client, xero_env, monkeypatch):
    c = http_client
    scopes = TOKEN_SCOPES + " accounting.journals.read"
    fake = FakeXero(scopes=scopes, journals=_many_journals())
    monkeypatch.setattr(xero_api, "default_http", fake)
    client_id, job_id = _new_client_and_job(c, "xero-journals@acme.test")
    _connect(c, client_id)
    c.post(f"/jobs/{job_id}/xero/pull")

    from app import storage, xero_pull
    job = storage.get_job(job_id)
    assert job["xero_pull"]["journals"]["status"] == "not_started"
    xc, _ = __import__("app.main", fromlist=["_xero_client"])._xero_client(client_id)

    ticks = iter(range(10_000))
    clock = lambda: next(ticks)  # each loop check advances 1s: a 1.5s budget = exactly one page per batch  # noqa: E731
    states = []
    for _ in range(6):
        job = storage.get_job(job_id)
        state = xero_pull.pull_journals_batch(job, xc, budget_seconds=1.5, clock=clock)
        states.append(state["status"])
        if state["status"] == "complete":
            break
    assert states[:2] == ["in_progress", "in_progress"] and states[-1] == "complete"

    job = storage.get_job(job_id)
    assert job["xero_pull"]["journals"]["kept"] == 130  # journals 121..250 only
    nominal = xero_pull.load_snapshot(job)["nominal_current"]
    assert len(nominal) == 260 and set(nominal["source_type"]) == {"Manual Journal"}
    assert nominal.groupby("account_code")["net"].sum().round(2).to_dict() == pytest.approx({"200": -sum(range(121, 251)), "610": sum(range(121, 251))})

    # scratch batch files were cleaned up once the frame was built
    assert job["xero_pull"]["journals"].get("batch_file_ids") is None

    # the pulled ledger now feeds the real pipeline
    from app import main
    assert len(main._load_canonical_data(job)["nominal_current"]) == 260


def test_ledger_refused_by_xero_degrades_to_unavailable_not_an_error(http_client, xero_env, monkeypatch):
    c = http_client
    fake = FakeXero(scopes=TOKEN_SCOPES + " accounting.journals.read", journals=_many_journals(), forbid_journals=True)
    monkeypatch.setattr(xero_api, "default_http", fake)
    client_id, job_id = _new_client_and_job(c, "xero-forbid@acme.test")
    _connect(c, client_id)
    c.post(f"/jobs/{job_id}/xero/pull")
    body = c.post(f"/jobs/{job_id}/xero/pull-journals").json()
    assert body["status"] == "unavailable" and "Account Transactions" in body["note"]


def test_repulling_replaces_the_previous_pull_and_its_files(http_client, xero_env, monkeypatch):
    c = http_client
    monkeypatch.setattr(xero_api, "default_http", FakeXero())
    client_id, job_id = _new_client_and_job(c, "xero-repull@acme.test")
    _connect(c, client_id)
    c.post(f"/jobs/{job_id}/xero/pull")
    from app import storage
    first_files = list(storage.get_job(job_id)["xero_pull"]["files"].values())
    c.post(f"/jobs/{job_id}/xero/pull")
    second_files = list(storage.get_job(job_id)["xero_pull"]["files"].values())
    assert set(first_files).isdisjoint(second_files)
    assert all(storage.load_file(f) is None for f in first_files)
    assert all(storage.load_file(f) is not None for f in second_files)


def test_raw_report_download_returns_exactly_what_xero_sent_and_404s_for_unknown(http_client, xero_env, monkeypatch):
    c = http_client
    fake = FakeXero()
    monkeypatch.setattr(xero_api, "default_http", fake)
    client_id, job_id = _new_client_and_job(c, "xero-raw@acme.test")
    _connect(c, client_id)
    c.post(f"/jobs/{job_id}/xero/pull")
    resp = c.get(f"/jobs/{job_id}/xero/raw/tb_current")
    assert resp.status_code == 200 and resp.json() == fake.cur_report
    assert c.get(f"/jobs/{job_id}/xero/raw/nonsense").status_code == 404
    assert 'href="/jobs/' in c.get(f"/jobs/{job_id}").text and "/xero/raw/tb_current" in c.get(f"/jobs/{job_id}").text
