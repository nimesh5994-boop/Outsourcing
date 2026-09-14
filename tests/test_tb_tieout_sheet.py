"""Excel-rendering tests for the TB Tie-Out sheet's Opening Balance
Disagreement table (tb_tieout._opening_balance_disagreements) - regression
coverage for a real gap found while adding the Comment column: since
Phase 2 moved this sheet onto its own custom formula-linked builder
(excel_builder.build_tb_tieout_sheet_formulas), that table stopped being
rendered at all (only the generic build_recon_sheet ever read
result.matched_detail) - the exception list was being silently computed
and then dropped.
"""
from openpyxl import Workbook

import pandas as pd

from app import tb_tieout
from app.data_sheets import write_data_sheets
from app.excel_builder import build_tb_tieout_sheet_formulas


def _tb(rows):
    df = pd.DataFrame(rows)
    df["balance"] = df["debit"] - df["credit"]
    return df


def test_opening_balance_disagreement_table_is_rendered_with_a_blank_comment_column():
    tb_current = _tb([{"account_code": "610", "account_name": "Trade Creditors", "account_type": "Current Liability", "debit": 0, "credit": 700}])
    tb_current_own_comparative = _tb([{"account_code": "610", "account_name": "Trade Creditors", "account_type": "Current Liability", "debit": 0, "credit": 500}])
    tb_comparative = _tb([{"account_code": "610", "account_name": "Trade Creditors", "account_type": "Current Liability", "debit": 0, "credit": 480}])

    result = tb_tieout.build_tieout(tb_current, tb_comparative, None, tb_current_own_comparative=tb_current_own_comparative)
    assert not result.matched_detail.empty  # sanity check on the fixture itself

    wb = Workbook()
    refs = write_data_sheets(wb, {"tb_current": tb_current, "tb_comparative": tb_comparative})
    ws = build_tb_tieout_sheet_formulas(wb, "Test Client", "Year ended 31 December 2025", "1", tb_current, result, refs)

    # find the matched_detail label and header row
    label_row = next(r for r in range(1, 20) if ws.cell(row=r, column=1).value == result.matched_detail_label.upper())
    header_row = label_row + 1
    headers = [ws.cell(row=header_row, column=c).value for c in range(1, 7)]
    assert headers == ["Account Code", "Account Name", "Per current TB's own comparative column",
                        "Per separately-uploaded comparative TB", "Diff", "Comment"]

    data_row = header_row + 1
    assert ws.cell(row=data_row, column=1).value == "610"
    assert ws.cell(row=data_row, column=5).value == -20.0  # Diff
    assert ws.cell(row=data_row, column=6).value == ""  # Comment - genuinely blank
    assert ws.cell(row=data_row, column=6).fill.fgColor.rgb == "00FFF9C4"  # the "type here" fill
